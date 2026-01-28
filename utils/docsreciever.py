import os
import io
import zipfile
import tempfile
from typing import List, Dict, Any
from PIL import Image
try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


class DocumentProcessor:
    def __init__(self, docs_folder: str = "docs"):
        self.docs_folder = docs_folder
        os.makedirs(self.docs_folder, exist_ok=True)

    def _ocr_image(self, image: Image.Image) -> str:
        if pytesseract is None:
            return ""
        try:
            return pytesseract.image_to_string(image).strip()
        except:
            return ""

    def process_pdf(self, pdf_path: str) -> List[Dict[str, Any]]:
        if pdfplumber is None:
            raise ImportError("pdfplumber not installed")
        
        extracted = []
        with pdfplumber.open(pdf_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_num = i + 1
                content = []
                
                text = page.extract_text() or ""
                if text.strip():
                    content.append(text)
                
                if page.images:
                    for img_obj in page.images:
                        try:
                            if all(k in img_obj for k in ['x0', 'top', 'x1', 'bottom']):
                                w = img_obj['x1'] - img_obj['x0']
                                h = img_obj['bottom'] - img_obj['top']
                                if w < 80 or h < 80:
                                    continue
                                bbox = (img_obj['x0'], img_obj['top'], img_obj['x1'], img_obj['bottom'])
                                cropped = page.crop(bbox)
                                pil_img = cropped.to_image(resolution=200).original
                                ocr = self._ocr_image(pil_img)
                                if ocr:
                                    content.append(f"[IMAGE OCR]\n{ocr}")
                        except:
                            pass
                
                if len(text.strip()) < 150:
                    try:
                        full_img = page.to_image(resolution=200).original
                        full_ocr = self._ocr_image(full_img)
                        if full_ocr:
                            content.append(f"[PAGE OCR]\n{full_ocr}")
                    except:
                        pass
                
                final = "\n".join(content)
                if final.strip():
                    extracted.append({
                        "page_num": page_num,
                        "text": final,
                        "source": os.path.basename(pdf_path),
                        "source_type": "PDF"
                    })
        return extracted

    def process_docx(self, docx_path: str) -> List[Dict[str, Any]]:
        if docx is None:
            raise ImportError("python-docx not installed")
        
        doc = docx.Document(docx_path)
        content = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    content.append(" | ".join(cells))
        
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref.lower():
                try:
                    img_data = rel.target_part.blob
                    img = Image.open(io.BytesIO(img_data))
                    ocr = self._ocr_image(img)
                    if ocr:
                        content.append(f"[DOCX IMAGE]\n{ocr}")
                except:
                    pass
        
        combined = "\n".join(content)
        if combined.strip():
            return [{
                "page_num": 1,
                "text": combined,
                "source": os.path.basename(docx_path),
                "source_type": "DOCX"
            }]
        return []

    def process_image(self, img_path: str) -> List[Dict[str, Any]]:
        try:
            img = Image.open(img_path)
            ocr = self._ocr_image(img)
            if ocr:
                return [{
                    "page_num": 1,
                    "text": ocr,
                    "source": os.path.basename(img_path),
                    "source_type": "IMAGE"
                }]
        except:
            pass
        return []

    def process_txt(self, txt_path: str) -> List[Dict[str, Any]]:
        with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            return [{
                "page_num": 1,
                "text": text,
                "source": os.path.basename(txt_path),
                "source_type": "TXT"
            }]
        return []

    def process_md(self, md_path: str) -> List[Dict[str, Any]]:
        with open(md_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            return [{
                "page_num": 1,
                "text": text,
                "source": os.path.basename(md_path),
                "source_type": "MD"
            }]
        return []

    def process_csv(self, csv_path: str) -> List[Dict[str, Any]]:
        if pd is None:
            raise ImportError("pandas not installed")
        
        df = pd.read_csv(csv_path)
        schema = f"Columns: {', '.join(df.columns.tolist())}"
        sample = df.head(10).to_string(index=False)
        full_data = df.to_string(index=False)
        
        return [{
            "page_num": 1,
            "text": f"{schema}\n\nSample Data:\n{sample}\n\nFull Data:\n{full_data}",
            "source": os.path.basename(csv_path),
            "source_type": "CSV",
            "is_structured": True,
            "dataframe": df
        }]

    def process_xlsx(self, xlsx_path: str) -> List[Dict[str, Any]]:
        if pd is None:
            raise ImportError("pandas not installed")
        
        xl = pd.ExcelFile(xlsx_path)
        results = []
        
        for sheet_name in xl.sheet_names:
            df = pd.read_excel(xl, sheet_name=sheet_name)
            schema = f"Sheet: {sheet_name}\nColumns: {', '.join(df.columns.tolist())}"
            sample = df.head(10).to_string(index=False)
            full_data = df.to_string(index=False)
            
            results.append({
                "page_num": xl.sheet_names.index(sheet_name) + 1,
                "text": f"{schema}\n\nSample Data:\n{sample}\n\nFull Data:\n{full_data}",
                "source": f"{os.path.basename(xlsx_path)}[{sheet_name}]",
                "source_type": "XLSX",
                "is_structured": True,
                "dataframe": df,
                "sheet_name": sheet_name
            })
        
        return results

    def process_pptx(self, pptx_path: str) -> List[Dict[str, Any]]:
        if Presentation is None:
            raise ImportError("python-pptx not installed")
        
        prs = Presentation(pptx_path)
        extracted = []
        
        for slide_num, slide in enumerate(prs.slides, 1):
            content = []
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    content.append(shape.text)
                
                if shape.shape_type == 13:
                    try:
                        img_stream = io.BytesIO(shape.image.blob)
                        img = Image.open(img_stream)
                        ocr = self._ocr_image(img)
                        if ocr:
                            content.append(f"[SLIDE IMAGE]\n{ocr}")
                    except:
                        pass
            
            if content:
                extracted.append({
                    "page_num": slide_num,
                    "text": "\n".join(content),
                    "source": os.path.basename(pptx_path),
                    "source_type": "PPTX"
                })
        
        return extracted

    def process_dockerfile(self, file_path: str) -> List[Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            return [{
                "page_num": 1,
                "text": text,
                "source": os.path.basename(file_path),
                "source_type": "DOCKERFILE"
            }]
        return []

    def process_zip(self, zip_path: str) -> List[Dict[str, Any]]:
        extracted = []
        with tempfile.TemporaryDirectory() as temp_dir:
            with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                zip_ref.extractall(temp_dir)
            
            for root, dirs, files in os.walk(temp_dir):
                for file in files:
                    file_path = os.path.join(root, file)
                    try:
                        docs = self.process_file(file_path)
                        for doc in docs:
                            doc["source"] = f"{os.path.basename(zip_path)}/{doc['source']}"
                        extracted.extend(docs)
                    except:
                        pass
        
        return extracted

    def process_file(self, file_path: str) -> List[Dict[str, Any]]:
        ext = os.path.splitext(file_path)[1].lower()
        basename = os.path.basename(file_path).lower()
        
        if ext == ".pdf":
            return self.process_pdf(file_path)
        elif ext == ".docx":
            return self.process_docx(file_path)
        elif ext == ".txt":
            return self.process_txt(file_path)
        elif ext == ".md":
            return self.process_md(file_path)
        elif ext == ".csv":
            return self.process_csv(file_path)
        elif ext == ".xlsx":
            return self.process_xlsx(file_path)
        elif ext == ".pptx":
            return self.process_pptx(file_path)
        elif ext == ".zip":
            return self.process_zip(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"]:
            return self.process_image(file_path)
        elif basename == "dockerfile" or ext == "":
            return self.process_dockerfile(file_path)
        else:
            try:
                return self.process_txt(file_path)
            except:
                raise ValueError(f"Unsupported: {ext}")

    def save_to_docs(self, documents: List[Dict[str, Any]], original_filename: str) -> str:
        base_name = os.path.splitext(original_filename)[0]
        output_path = os.path.join(self.docs_folder, f"{base_name}.txt")
        
        content_parts = []
        for doc in documents:
            header = f"=== {doc['source']} | Page {doc['page_num']} | Type: {doc['source_type']} ==="
            content_parts.append(f"{header}\n{doc['text']}\n")
        
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(content_parts))
        
        return output_path

    def process_and_save(self, file_path: str) -> Dict[str, Any]:
        documents = self.process_file(file_path)
        saved_path = self.save_to_docs(documents, os.path.basename(file_path))
        
        is_structured = any(doc.get("is_structured", False) for doc in documents)
        
        return {
            "original_file": file_path,
            "saved_path": saved_path,
            "documents": documents,
            "is_structured": is_structured,
            "doc_count": len(documents)
        }
